from django.core.management.base import BaseCommand, CommandError
from django.db import models
from testcase.models import TestcaseType, MessageType
import os, sys, json
import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl import utils

class Command(BaseCommand):
    help = 'Import testcase data from Excel file to database'

    def __init__(self) -> None:
        super().__init__()

    def add_arguments(self, parser):
        # positional argument
        parser.add_argument('excel_file', type=str, help='Path to the Excel file containing testcases')
        parser.add_argument('--project_path', type=str, help='Path to the C++ test project (optional)')

    def handle(self, *args, **options):
        excel_file = options['excel_file']
        project_path = options.get('project_path', '')

        if not os.path.isfile(excel_file):
            raise CommandError(f"Excel file '{excel_file}' does not exist")

        try:
            self.stdout.write(f"Importing testcases from {excel_file}...")

            # Load Excel workbook
            wb = openpyxl.load_workbook(excel_file, read_only=True)
            ws = wb["testcase"]

            # Import testcases
            imported_count = self.import_testcases(ws, project_path)

            wb.close()

            self.stdout.write(
                self.style.SUCCESS(
                    f"Successfully imported {imported_count} testcases from Excel file"
                )
            )

        except Exception as e:
            raise CommandError(f"Error importing testcases: {str(e)}")

    def import_testcases(self, ws: Worksheet, project_path: str) -> int:
        """Import testcases from Excel worksheet"""
        imported_count = 0

        # Skip header row
        for row in range(2, ws.max_row + 1):
            try:
                # Read data from Excel row
                testcase_id = ws.cell(row, 1).value
                name = ws.cell(row, 2).value
                tc_type = ws.cell(row, 3).value
                condition = ws.cell(row, 4).value or ""
                expected = ws.cell(row, 5).value or ""
                pics = ws.cell(row, 6).value or ""
                pixit = ws.cell(row, 7).value or ""
                msg_type_id = ws.cell(row, 8).value
                message_sets = ws.cell(row, 9).value or ""
                result = ws.cell(row, 10).value or "none"

                # Skip if name is empty
                if not name or str(name).strip() == "":
                    continue

                # Validate testcase type
                if tc_type not in ['AC', 'DC', 'CMN']:
                    self.stdout.write(
                        self.style.WARNING(
                            f"Skipping row {row}: Invalid type '{tc_type}' for testcase '{name}'"
                        )
                    )
                    continue

                # Handle message type
                msg_type = None
                if msg_type_id:
                    try:
                        # Try to get existing message type or create new one
                        msg_type, _ = MessageType.objects.get_or_create(
                            name=f"MessageType_{msg_type_id}"
                        )
                    except Exception as e:
                        self.stdout.write(
                            self.style.WARNING(
                                f"Could not create message type for {name}: {str(e)}"
                            )
                        )

                # Prepare path
                path = ""
                if project_path:
                    path = os.path.join(project_path, 'build/TestSuite/runTC')

                # Convert result to status
                status_map = {
                    'none': 'n',
                    'pass': 'p',
                    'inconclusion': 'i',
                    'fail': 'f',
                    'error': 'e'
                }
                status = status_map.get(str(result).lower(), 'n')

                # Clean up condition and expected (remove extra newlines)
                condition = str(condition).strip()
                expected = str(expected).strip()
                pics = str(pics).strip()
                pixit = str(pixit).strip()

                # Try to get existing testcase or create new one
                try:
                    tc = TestcaseType.objects.get(name=name)
                    # Update existing testcase
                    tc.type = tc_type
                    tc.status = status
                    tc.condition = condition
                    tc.expected = expected
                    tc.pics = pics
                    tc.pixit = pixit
                    tc.path = path
                    tc.msg_type = msg_type
                    tc.build_status = "NA"  # Not available since we're not building C++
                    tc.save()
                    self.stdout.write(f"Updated existing testcase: {name}")
                except TestcaseType.DoesNotExist:
                    # Create new testcase
                    tc = TestcaseType.objects.create(
                        name=name,
                        type=tc_type,
                        status=status,
                        build_status="NA",
                        condition=condition,
                        expected=expected,
                        pics=pics,
                        pixit=pixit,
                        path=path,
                        msg_type=msg_type
                    )
                    self.stdout.write(f"Created new testcase: {name}")

                imported_count += 1

            except Exception as e:
                self.stdout.write(
                    self.style.ERROR(
                        f"Error processing row {row}: {str(e)}"
                    )
                )
                continue

        return imported_count

